"""Scrape books.toscrape.com (a sandbox site built for scraping practice) into clean CSV + Excel."""
import argparse, csv, time
from urllib.parse import urljoin
import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE = "https://books.toscrape.com/catalogue/page-1.html"
RATING = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

def session(retries=4):
    s = requests.Session()
    r = Retry(total=retries, backoff_factor=0.5, status_forcelist=[429, 500, 502, 503, 504])
    s.mount("https://", HTTPAdapter(max_retries=r)); s.mount("http://", HTTPAdapter(max_retries=r))
    s.headers["User-Agent"] = "portfolio-demo-scraper/1.0"
    return s

def parse_page(html, url):
    soup = BeautifulSoup(html, "html.parser"); rows = []
    for a in soup.select("article.product_pod"):
        link = a.h3.a
        rows.append({"title": link["title"].strip(),
                     "price_gbp": float(a.select_one(".price_color").text.strip().lstrip("Â£")),
                     "rating": RATING.get(a.select_one(".star-rating")["class"][1], 0),
                     "in_stock": "In stock" in a.select_one(".availability").text,
                     "url": urljoin(url, link["href"])})
    nxt = soup.select_one("li.next a")
    return rows, (urljoin(url, nxt["href"]) if nxt else None)

def dedupe(rows):
    seen, out = set(), []
    for r in rows:
        if r["url"] not in seen: seen.add(r["url"]); out.append(r)
    return out

def scrape(start=BASE, max_pages=None, delay=0.3, sess=None):
    sess = sess or session(); url, rows, n = start, [], 0
    while url and (max_pages is None or n < max_pages):
        resp = sess.get(url, timeout=15); resp.raise_for_status()
        page, url = parse_page(resp.content.decode("utf-8"), resp.url); rows += page; n += 1
        print(f"page {n}: {len(page)} items"); time.sleep(delay)
    return rows, n

def export(rows, stem):
    with open(f"{stem}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=rows[0].keys()); w.writeheader(); w.writerows(rows)
    wb = Workbook(); ws = wb.active; ws.title = "Books"; ws.append(list(rows[0].keys()))
    for r in rows: ws.append(list(r.values()))
    for c in ws[1]: c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    for col, w in zip("ABCDE", (60, 11, 8, 10, 70)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions; wb.save(f"{stem}.xlsx")

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--pages", type=int, default=None, help="max pages (default: all 50)")
    ap.add_argument("--out", default="output/books")
    a = ap.parse_args(argv)
    raw, n = scrape(max_pages=a.pages); rows = dedupe(raw)
    rows.sort(key=lambda r: (-r["rating"], r["price_gbp"]))
    export(rows, a.out)
    print(f"\nPages: {n} | scraped: {len(raw)} | after dedupe: {len(rows)} -> {a.out}.csv / .xlsx")
    return rows

if __name__ == "__main__":
    main()
